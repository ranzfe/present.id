import streamlit as st
import pandas as pd
from pypdf import PdfReader
from pptx import Presentation
import json
import io
import openai

# --- CONFIGURATION & PAGE SETUP ---
st.set_page_config(page_title="present.id - AI Presentation Generator", page_icon="📊", layout="wide")

st.title("📊 present.id")
st.subheader("Ubah Dokumen (PDF/Excel) Menjadi Presentasi PPTX dalam Sekejap")

# Sidebar untuk API Key
with st.sidebar:
    st.header("Pengaturan")
    api_key = st.text_input("Masukkan OpenAI API Key:", type="password")

# --- HELPER FUNCTIONS ---
def extract_text_from_pdf(pdf_file):
    reader = PdfReader(pdf_file)
    text = ""
    for page in reader.pages:
        text += page.extract_text() or ""
    return text

def extract_data_from_excel(excel_file):
    xls = pd.ExcelFile(excel_file)
    summary_text = ""
    for sheet in xls.sheet_names:
        df = pd.read_excel(excel_file, sheet_name=sheet)
        summary_text += f"\n--- Sheet: {sheet} ---\n"
        summary_text += df.head(10).to_string() # Mengambil sampel 10 baris pertama
    return summary_text

def generate_slides_json(content, openai_api_key):
    client = openai.OpenAI(api_key=openai_api_key)
    
    prompt = f"""
    Kamu adalah asisten profesional pembuat presentasi.
    Analisis dokumen/data berikut dan buatlah struktur presentasi 5-7 slide.
    
    Kembalikan respons HANYA dalam format JSON dengan struktur seperti ini:
    [
        {{
            "title": "Judul Slide 1",
            "subtitle": "Subjudul atau deskripsi singkat",
            "bullets": ["Poin 1", "Poin 2", "Poin 3"]
        }}
    ]

    Dokumen/Data:
    {content[:4000]}  # Mencegah over-token
    """
    
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "system", "content": "You output JSON only."},
                  {"role": "user", "content": prompt}],
        response_format={"type": "json_object"}
    )
    
    return json.loads(response.choices[0].message.content)

def create_pptx(slides_data):
    prs = Presentation()
    
    for slide_info in slides_data.get("slides", slides_data if isinstance(slides_data, list) else []):
        slide_layout = prs.slide_layouts[1] # Layout Title & Content
        slide = prs.slides.add_slide(slide_layout)
        
        # Set Title
        title = slide.shapes.title
        title.text = slide_info.get("title", "")
        
        # Set Content (Subtitle & Bullets)
        tf = slide.placeholders[1].text_frame
        tf.text = slide_info.get("subtitle", "")
        
        for bullet in slide_info.get("bullets", []):
            p = tf.add_paragraph()
            p.text = bullet
            p.level = 1
            
    # Save to buffer
    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    pptx_io.seek(0)
    return pptx_io

# --- MAIN APP LOGIC ---
uploaded_file = st.file_uploader("Upload file PDF atau Excel kamu di sini", type=["pdf", "xlsx", "xls"])

if uploaded_file and api_key:
    if st.button("🚀 Buat Presentasi"):
        with st.spinner("Menganalisis dokumen dan menggenerate slide..."):
            extracted_text = ""
            
            # 1. Ekstraksi Dokumen
            if uploaded_file.name.endswith('.pdf'):
                extracted_text = extract_text_from_pdf(uploaded_file)
            elif uploaded_file.name.endswith(('.xlsx', '.xls')):
                extracted_text = extract_data_from_excel(uploaded_file)
                
            # 2. Olah dengan AI
            slides_data = generate_slides_json(extracted_text, api_key)
            
            # 3. Generate PPTX
            pptx_file = create_pptx(slides_data)
            
            st.success("Presentasi berhasil dibuat!")
            
            # 4. Download Button
            st.download_button(
                label="📥 Download File PPTX",
                data=pptx_file,
                file_name=f"{uploaded_file.name.split('.')[0]}_presentation.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
            )
elif uploaded_file and not api_key:
    st.warning("Silakan masukkan OpenAI API Key kamu di sidebar terlebih dahulu.")
